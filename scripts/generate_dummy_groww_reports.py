from pathlib import Path

from openpyxl import Workbook


OUTPUT = Path("local_test_files")
OUTPUT.mkdir(parents=True, exist_ok=True)


def create_shares_report():
    wb = Workbook()
    ws = wb.active

    ws.title = "Tax P&L"

    ws["A1"] = (
        "Groww Shares Tax Statement "
        "from 01 Apr 2025 to 31 Mar 2026"
    )

    ws["A3"] = "Name"
    ws["B3"] = "Dummy Client"

    ws["A4"] = "Unique Client Code"
    ws["B4"] = "9061285890"

    # -----------------------------
    # SUMMARY
    # -----------------------------

    ws["A6"] = "Intraday Summary"

    ws["A7"] = "Intraday Turnover"
    ws["B7"] = 24000

    ws["A8"] = "Intraday P&L"
    ws["B8"] = 6000

    ws["A10"] = "Capital Gains Summary"

    ws["A11"] = "STCG"
    ws["B11"] = 45000

    ws["A12"] = "LTCG"
    ws["B12"] = 120000

    ws["A13"] = "Sell Value"
    ws["B13"] = 600000

    # -----------------------------
    # INTRADAY TRADE-LEVEL DATA
    # -----------------------------

    ws["A15"] = "Intraday Trades"

    ws["A16"] = "Trade"
    ws["B16"] = "Realised P&L"

    trades = [
        ("Trade 1", 10000),
        ("Trade 2", -7000),
        ("Trade 3", 5000),
        ("Trade 4", -2000),
    ]

    for row, (trade, pnl) in enumerate(
        trades,
        start=17,
    ):
        ws.cell(row=row, column=1).value = trade
        ws.cell(row=row, column=2).value = pnl

    path = OUTPUT / "dummy_groww_shares.xlsx"

    wb.save(path)

    print(f"Created: {path}")


def create_mutual_fund_report():
    wb = Workbook()
    ws = wb.active

    ws.title = "Mutual Funds"

    ws["A1"] = (
        "Groww Mutual Funds Tax Statement "
        "from 01 Apr 2025 to 31 Mar 2026"
    )

    ws["A3"] = "Name"
    ws["B3"] = "Dummy Client"

    ws["A4"] = "Unique Client Code"
    ws["B4"] = "GROWW-DUMMY-001"

    # -----------------------------
    # CAPITAL GAINS SUMMARY
    # -----------------------------

    ws["A6"] = "Capital Gains Summary"

    ws["A7"] = "Taxable STCG"
    ws["B7"] = 3500

    ws["A8"] = "Taxable LTCG"
    ws["B8"] = 8200

    # -----------------------------
    # REDEMPTION TRANSACTIONS
    # -----------------------------

    ws["A10"] = "Redemption Transactions"

    ws["A11"] = "Scheme"
    ws["B11"] = "Redeem Price"
    ws["C11"] = "Matched Quantity"
    ws["D11"] = "Tax Category"

    transactions = [
        (
            "Dummy Equity Fund A",
            125.50,
            100,
            "STCG",
        ),
        (
            "Dummy Equity Fund B",
            210.00,
            50,
            "LTCG",
        ),
        (
            "Dummy Equity Fund C",
            85.25,
            200,
            "STCG",
        ),
    ]

    for row, transaction in enumerate(
        transactions,
        start=12,
    ):
        for column, value in enumerate(
            transaction,
            start=1,
        ):
            ws.cell(
                row=row,
                column=column,
            ).value = value

    path = (
        OUTPUT
        / "dummy_groww_mutual_funds.xlsx"
    )

    wb.save(path)

    print(f"Created: {path}")


if __name__ == "__main__":
    create_shares_report()
    create_mutual_fund_report()