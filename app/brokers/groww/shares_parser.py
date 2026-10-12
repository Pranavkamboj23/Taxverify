from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import re

from openpyxl import load_workbook

from app.brokers.groww.parser import (
    GrowwParseError,
    SourceValue,
    find_exact,
    normalize_text,
    to_decimal,
)


@dataclass
class GrowwSharesParseResult:
    report_type: str

    client_name: str
    client_code: str

    period_start: object
    period_end: object

    intraday_turnover: SourceValue
    intraday_net_pnl: SourceValue

    stcg: SourceValue
    ltcg: SourceValue

    sale_consideration: Decimal

    calculated_intraday_turnover: Decimal
    calculated_intraday_net_pnl: Decimal

    calculated_stcg: Decimal
    calculated_ltcg: Decimal


# =========================================================
# FIND REAL GROWW STOCKS SHEET
# =========================================================

def find_stocks_sheet(workbook):
    """
    Real Groww Stocks exports do not necessarily
    use a fixed sheet name.

    The sample provided uses:
        Sheet1

    Therefore identify it using the report title.
    """

    for sheet in workbook.worksheets:

        for row in range(
            1,
            min(sheet.max_row, 15) + 1,
        ):

            value = normalize_text(
                sheet.cell(
                    row=row,
                    column=1,
                ).value
            )

            if (
                "Capital Gains Statement for stocks"
                in value
            ):
                return sheet

    raise GrowwParseError(
        "Could not locate a supported "
        "Groww Stocks Capital Gains report."
    )


# =========================================================
# IDENTITY
# =========================================================

def extract_identity(sheet):

    client_name = None
    client_code = None

    for row in range(
        1,
        min(sheet.max_row, 15) + 1,
    ):

        label = normalize_text(
            sheet.cell(
                row=row,
                column=1,
            ).value
        )

        value = normalize_text(
            sheet.cell(
                row=row,
                column=2,
            ).value
        )

        if label.casefold() == "name":
            client_name = value

        elif (
            label.casefold()
            == "unique client code"
        ):
            client_code = value

    if not client_name:
        raise GrowwParseError(
            "Groww Stocks client name "
            "was not found."
        )

    if not client_code:
        raise GrowwParseError(
            "Groww Stocks Unique Client "
            "Code was not found."
        )

    return (
        client_name,
        client_code,
    )


# =========================================================
# PERIOD
# =========================================================

def extract_stocks_period(
    title: str,
):
    """
    Actual Groww example:

    Capital Gains Statement for stocks
    from 01-04-2025 To 31-03-2026
    """

    match = re.search(
        r"from\s+"
        r"(\d{2}-\d{2}-\d{4})"
        r"\s+to\s+"
        r"(\d{2}-\d{2}-\d{4})",
        title,
        re.IGNORECASE,
    )

    if not match:
        raise GrowwParseError(
            "Could not determine Groww Stocks "
            "report period."
        )

    period_start = datetime.strptime(
        match.group(1),
        "%d-%m-%Y",
    ).date()

    period_end = datetime.strptime(
        match.group(2),
        "%d-%m-%Y",
    ).date()

    return (
        period_start,
        period_end,
    )


# =========================================================
# SOURCE VALUE
# =========================================================

def source_value_from_label(
    sheet,
    label: str,
) -> SourceValue:

    cell = find_exact(
        sheet,
        label,
    )

    if not cell:
        raise GrowwParseError(
            f"Required Groww Stocks field "
            f"not found: {label}"
        )

    value_cell = sheet.cell(
        row=cell.row,
        column=cell.column + 1,
    )

    return SourceValue(
        value=to_decimal(
            value_cell.value
        ),
        sheet=sheet.title,
        cell=value_cell.coordinate,
        label=label,
    )


# =========================================================
# TRANSACTION SECTION
# =========================================================

def read_trade_section(
    sheet,
    section_name: str,
):
    """
    Groww actual structure:

    Short Term trades
    Stock name | ISIN | Quantity | ... |
    Sell Value | Realised P&L | Remark
    """

    heading = find_exact(
        sheet,
        section_name,
    )

    if not heading:
        return []

    header_row = heading.row + 1
    data_row = header_row + 1

    sell_value_column = None
    pnl_column = None

    for column in range(
        1,
        sheet.max_column + 1,
    ):

        header = normalize_text(
            sheet.cell(
                row=header_row,
                column=column,
            ).value
        ).casefold()

        if header == "sell value":
            sell_value_column = column

        elif (
            header
            in {
                "realised p&l",
                "realized p&l",
            }
        ):
            pnl_column = column

    if pnl_column is None:
        raise GrowwParseError(
            f"Realised P&L column was not "
            f"found in '{section_name}'."
        )

    rows = []

    row = data_row

    while row <= sheet.max_row:

        stock_name = normalize_text(
            sheet.cell(
                row=row,
                column=1,
            ).value
        )

        if not stock_name:
            break

        pnl_value = sheet.cell(
            row=row,
            column=pnl_column,
        ).value

        if pnl_value is None:
            break

        sell_value = Decimal("0")

        if sell_value_column:
            raw_sell = sheet.cell(
                row=row,
                column=sell_value_column,
            ).value

            if raw_sell is not None:
                sell_value = to_decimal(
                    raw_sell
                )

        rows.append(
            {
                "sell_value": sell_value,
                "pnl": to_decimal(
                    pnl_value
                ),
            }
        )

        row += 1

    return rows


# =========================================================
# PARSER
# =========================================================

def parse_groww_shares(
    filepath: str | Path,
) -> GrowwSharesParseResult:

    workbook = load_workbook(
        filename=filepath,
        data_only=True,
        read_only=False,
    )

    sheet = find_stocks_sheet(
        workbook
    )

    title = normalize_text(
        sheet["A4"].value
    )

    if (
        "Capital Gains Statement for stocks"
        not in title
    ):
        raise GrowwParseError(
            "This is not a supported "
            "Groww Stocks report."
        )

    (
        client_name,
        client_code,
    ) = extract_identity(
        sheet
    )

    (
        period_start,
        period_end,
    ) = extract_stocks_period(
        title
    )

    # -----------------------------------------------------
    # SUMMARY VALUES
    # -----------------------------------------------------

    intraday_net_pnl = (
        source_value_from_label(
            sheet,
            "Intraday P&L",
        )
    )

    stcg = source_value_from_label(
        sheet,
        "Short Term P&L",
    )

    ltcg = source_value_from_label(
        sheet,
        "Long Term P&L",
    )

    # Turnover section uses label "Intraday"
    turnover_heading = find_exact(
        sheet,
        "Turnover",
    )

    if not turnover_heading:
        raise GrowwParseError(
            "Groww Stocks Turnover "
            "section was not found."
        )

    intraday_turnover_cell = None

    for row in range(
        turnover_heading.row + 1,
        min(
            turnover_heading.row + 10,
            sheet.max_row,
        ) + 1,
    ):

        label = normalize_text(
            sheet.cell(
                row=row,
                column=turnover_heading.column,
            ).value
        )

        if label.casefold() == "intraday":

            intraday_turnover_cell = (
                sheet.cell(
                    row=row,
                    column=(
                        turnover_heading.column
                        + 1
                    ),
                )
            )

            break

    if intraday_turnover_cell is None:
        raise GrowwParseError(
            "Groww Stocks Intraday "
            "Turnover was not found."
        )

    intraday_turnover = SourceValue(
        value=to_decimal(
            intraday_turnover_cell.value
        ),
        sheet=sheet.title,
        cell=(
            intraday_turnover_cell.coordinate
        ),
        label="Intraday Turnover",
    )

    # -----------------------------------------------------
    # TRANSACTION LEVEL
    # -----------------------------------------------------

    intraday_trades = read_trade_section(
        sheet,
        "Intraday trades",
    )

    short_term_trades = read_trade_section(
        sheet,
        "Short Term trades",
    )

    long_term_trades = read_trade_section(
        sheet,
        "Long Term trades",
    )

    if not intraday_trades:
        calculated_intraday_turnover = (
            Decimal("0")
        )

        calculated_intraday_net_pnl = (
            Decimal("0")
        )

    else:

        calculated_intraday_turnover = sum(
            (
                abs(
                    trade["pnl"]
                )
                for trade
                in intraday_trades
            ),
            Decimal("0"),
        )

        calculated_intraday_net_pnl = sum(
            (
                trade["pnl"]
                for trade
                in intraday_trades
            ),
            Decimal("0"),
        )

    # -----------------------------------------------------
    # CAPITAL GAINS
    # -----------------------------------------------------

    calculated_stcg = sum(
        (
            trade["pnl"]
            for trade
            in short_term_trades
        ),
        Decimal("0"),
    )

    calculated_ltcg = sum(
        (
            trade["pnl"]
            for trade
            in long_term_trades
        ),
        Decimal("0"),
    )

    # -----------------------------------------------------
    # SALE CONSIDERATION
    #
    # Investment/delivery shares only.
    # Intraday Sell Value is NOT capital-gain
    # sale consideration.
    # -----------------------------------------------------

    sale_consideration = sum(
        (
            trade["sell_value"]
            for trade
            in (
                short_term_trades
                + long_term_trades
            )
        ),
        Decimal("0"),
    )

    return GrowwSharesParseResult(
        report_type="SHARES",

        client_name=client_name,
        client_code=client_code,

        period_start=period_start,
        period_end=period_end,

        intraday_turnover=(
            intraday_turnover
        ),

        intraday_net_pnl=(
            intraday_net_pnl
        ),

        stcg=stcg,
        ltcg=ltcg,

        sale_consideration=(
            sale_consideration
        ),

        calculated_intraday_turnover=(
            calculated_intraday_turnover
        ),

        calculated_intraday_net_pnl=(
            calculated_intraday_net_pnl
        ),

        calculated_stcg=(
            calculated_stcg
        ),

        calculated_ltcg=(
            calculated_ltcg
        ),
    )