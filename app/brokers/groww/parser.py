import re

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet


# =========================================================
# EXCEPTIONS
# =========================================================


class GrowwParseError(Exception):
    pass


# =========================================================
# DATA TYPES
# =========================================================


class SourceValue:
    def __init__(
        self,
        value: Decimal,
        sheet: str,
        cell: str,
        label: str,
    ):
        self.value = value
        self.sheet = sheet
        self.cell = cell
        self.label = label


class GrowwParseResult:
    def __init__(
        self,
        report_type: str,
        client_name: str,
        client_code: str,
        period_start,
        period_end,
        realised_pnl: SourceValue,
        charges: SourceValue,
        reported_turnover,
        calculated_turnover: Decimal,
        turnover_source: str,
        trade_level_pnl: Decimal,
        final_net_pnl: Decimal,
    ):
        self.report_type = report_type

        self.client_name = client_name
        self.client_code = client_code

        self.period_start = period_start
        self.period_end = period_end

        self.realised_pnl = realised_pnl
        self.charges = charges

        self.reported_turnover = (
            reported_turnover
        )

        self.calculated_turnover = (
            calculated_turnover
        )

        self.turnover_source = (
            turnover_source
        )

        self.trade_level_pnl = (
            trade_level_pnl
        )

        self.final_net_pnl = (
            final_net_pnl
        )


# =========================================================
# GENERAL HELPERS
# =========================================================


def normalize_text(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def to_decimal(value) -> Decimal:
    if value is None:
        return Decimal("0")

    if isinstance(
        value,
        Decimal,
    ):
        return value

    if isinstance(
        value,
        (int, float),
    ):
        return Decimal(
            str(value)
        )

    text = (
        str(value)
        .strip()
        .replace(",", "")
        .replace("₹", "")
    )

    if not text:
        return Decimal("0")

    try:
        return Decimal(text)

    except (
        InvalidOperation,
        ValueError,
        TypeError,
    ):
        raise GrowwParseError(
            f"Could not convert value "
            f"to Decimal: {value}"
        )


def try_decimal(value):
    if value is None:
        return None

    try:
        return to_decimal(
            value
        )

    except GrowwParseError:
        return None


# =========================================================
# FIND EXACT TEXT
# =========================================================


def find_exact(
    sheet: Worksheet,
    text: str,
):
    target = (
        text
        .strip()
        .casefold()
    )

    for row in sheet.iter_rows():

        for cell in row:

            value = normalize_text(
                cell.value
            )

            if (
                value.casefold()
                == target
            ):
                return cell

    return None


# =========================================================
# FIND REPORT TITLE SHEET
# =========================================================


def find_title_sheet(
    workbook,
    report_type: str,
):
    """
    Supports the actual Groww formats
    encountered so far.

    F&O examples:

    P&L Statement for Futures & Options
    from 01 Apr 2025 to 31 Mar 2026

    Tax Statement for F&O
    from 01 Apr 2025 to 04 Mar 2026

    Commodity example:

    Tax Statement for Commodities
    from 01 Apr 2025 to 31 Mar 2026
    """

    report_type = (
        report_type.upper()
    )

    for sheet in workbook.worksheets:

        max_row = min(
            sheet.max_row,
            25,
        )

        max_column = min(
            sheet.max_column,
            8,
        )

        for row in range(
            1,
            max_row + 1,
        ):

            for column in range(
                1,
                max_column + 1,
            ):

                value = normalize_text(
                    sheet.cell(
                        row=row,
                        column=column,
                    ).value
                )

                if not value:
                    continue

                lowered = (
                    value.casefold()
                )

                # -----------------------------------------
                # F&O
                # -----------------------------------------

                if report_type == "FNO":

                    if (
                        "p&l statement for futures & options"
                        in lowered
                        or
                        "tax statement for f&o"
                        in lowered
                        or
                        "tax statement for futures & options"
                        in lowered
                    ):
                        return (
                            sheet,
                            value,
                        )

                # -----------------------------------------
                # COMMODITY
                # -----------------------------------------

                elif (
                    report_type
                    == "COMMODITY"
                ):

                    if (
                        "tax statement for commodities"
                        in lowered
                        or
                        "p&l statement for commodities"
                        in lowered
                    ):
                        return (
                            sheet,
                            value,
                        )

    raise GrowwParseError(
        f"Could not locate a supported "
        f"Groww {report_type} report."
    )


# =========================================================
# EXTRACT PERIOD
# =========================================================


def extract_period(
    title: str,
):
    """
    Supports:

    01 Apr 2025 to 31 Mar 2026

    and

    01-04-2025 to 31-03-2026
    """

    # -----------------------------------------------------
    # DD Mon YYYY
    # -----------------------------------------------------

    match = re.search(
        r"from\s+"
        r"(\d{1,2}\s+[A-Za-z]{3}\s+\d{4})"
        r"\s+to\s+"
        r"(\d{1,2}\s+[A-Za-z]{3}\s+\d{4})",
        title,
        re.IGNORECASE,
    )

    if match:

        try:
            start = datetime.strptime(
                match.group(1),
                "%d %b %Y",
            ).date()

            end = datetime.strptime(
                match.group(2),
                "%d %b %Y",
            ).date()

            return (
                start,
                end,
            )

        except ValueError:
            pass

    # -----------------------------------------------------
    # DD-MM-YYYY
    # -----------------------------------------------------

    match = re.search(
        r"from\s+"
        r"(\d{2}-\d{2}-\d{4})"
        r"\s+to\s+"
        r"(\d{2}-\d{2}-\d{4})",
        title,
        re.IGNORECASE,
    )

    if match:

        try:
            start = datetime.strptime(
                match.group(1),
                "%d-%m-%Y",
            ).date()

            end = datetime.strptime(
                match.group(2),
                "%d-%m-%Y",
            ).date()

            return (
                start,
                end,
            )

        except ValueError:
            pass

    raise GrowwParseError(
        "Could not determine "
        "Groww report period."
    )


# =========================================================
# EXTRACT IDENTITY
# =========================================================


def find_label_value(
    sheet: Worksheet,
    label: str,
):
    """
    Search the top part of the workbook for
    a label and return the first non-empty
    cell to its right.
    """

    target = (
        label
        .strip()
        .casefold()
    )

    max_row = min(
        sheet.max_row,
        25,
    )

    max_column = min(
        sheet.max_column,
        10,
    )

    for row in range(
        1,
        max_row + 1,
    ):

        for column in range(
            1,
            max_column + 1,
        ):

            value = normalize_text(
                sheet.cell(
                    row=row,
                    column=column,
                ).value
            )

            if (
                value.casefold()
                != target
            ):
                continue

            for next_column in range(
                column + 1,
                max_column + 1,
            ):

                result = normalize_text(
                    sheet.cell(
                        row=row,
                        column=next_column,
                    ).value
                )

                if result:
                    return result

    return None


def extract_identity(
    sheet: Worksheet,
):
    client_name = (
        find_label_value(
            sheet,
            "Name",
        )
    )

    client_code = (
        find_label_value(
            sheet,
            "Unique Client Code",
        )
    )

    if not client_name:
        raise GrowwParseError(
            "Groww client name "
            "was not found."
        )

    if not client_code:
        raise GrowwParseError(
            "Groww Unique Client Code "
            "was not found."
        )

    return (
        client_name,
        client_code,
    )


# =========================================================
# SECTION TOTAL
# =========================================================


def find_section_total(
    sheet: Worksheet,
    section_name: str,
) -> SourceValue:

    section_cell = find_exact(
        sheet,
        section_name,
    )

    if not section_cell:
        raise GrowwParseError(
            f"Groww section not found: "
            f"{section_name}"
        )

    start_row = (
        section_cell.row + 1
    )

    end_row = min(
        start_row + 25,
        sheet.max_row,
    )

    for row in range(
        start_row,
        end_row + 1,
    ):

        label_cell = sheet.cell(
            row=row,
            column=section_cell.column,
        )

        label = normalize_text(
            label_cell.value
        )

        # Stop if we've clearly reached
        # another major section.
        if (
            row > start_row
            and label.casefold()
            in {
                "turnover",
                "charges",
                "realised trades (trade level)",
                "realized trades (trade level)",
            }
        ):
            break

        if (
            label.casefold()
            == "total"
        ):

            value_cell = sheet.cell(
                row=row,
                column=(
                    section_cell.column
                    + 1
                ),
            )

            return SourceValue(
                value=to_decimal(
                    value_cell.value
                ),
                sheet=sheet.title,
                cell=(
                    value_cell.coordinate
                ),
                label=(
                    f"{section_name} Total"
                ),
            )

    raise GrowwParseError(
        f"Total was not found for "
        f"Groww section: "
        f"{section_name}"
    )


# =========================================================
# REPORTED TURNOVER
# =========================================================


def find_optional_turnover(
    sheet: Worksheet,
):
    """
    Groww versions encountered:

    Turnover
        Futures    X
        Options    Y

    Some older F&O files do not include
    this summary section at all.

    When present:
        Futures + Options

    becomes broker-reported turnover.
    """

    turnover_cell = find_exact(
        sheet,
        "Turnover",
    )

    if not turnover_cell:
        return None

    start_row = (
        turnover_cell.row + 1
    )

    end_row = min(
        start_row + 15,
        sheet.max_row,
    )

    total = Decimal("0")
    found = False

    source_cells = []

    for row in range(
        start_row,
        end_row + 1,
    ):

        label = normalize_text(
            sheet.cell(
                row=row,
                column=turnover_cell.column,
            ).value
        )

        label_lower = (
            label.casefold()
        )

        # -----------------------------------------
        # Explicit total if Groww introduces one
        # -----------------------------------------

        if (
            label_lower == "total"
        ):

            value_cell = sheet.cell(
                row=row,
                column=(
                    turnover_cell.column
                    + 1
                ),
            )

            explicit_total = (
                try_decimal(
                    value_cell.value
                )
            )

            if explicit_total is not None:

                return SourceValue(
                    value=explicit_total,
                    sheet=sheet.title,
                    cell=(
                        value_cell.coordinate
                    ),
                    label=(
                        "Turnover Total"
                    ),
                )

        # -----------------------------------------
        # Futures / Options
        # -----------------------------------------

        if (
            label_lower
            in {
                "futures",
                "options",
            }
        ):

            value_cell = sheet.cell(
                row=row,
                column=(
                    turnover_cell.column
                    + 1
                ),
            )

            value = try_decimal(
                value_cell.value
            )

            if value is not None:

                total += value

                source_cells.append(
                    value_cell.coordinate
                )

                found = True

                continue

        # -----------------------------------------
        # Stop at next section
        # -----------------------------------------

        if (
            found
            and label_lower
            in {
                "charges",
                "realised trades (trade level)",
                "realized trades (trade level)",
            }
        ):
            break

    if not found:
        return None

    return SourceValue(
        value=total,
        sheet=sheet.title,
        cell=(
            ", ".join(
                source_cells
            )
        ),
        label=(
            "Turnover "
            "(Futures + Options)"
        ),
    )


# =========================================================
# TRADE-LEVEL RECONCILIATION
# =========================================================


def calculate_trade_level_turnover(
    sheet: Worksheet,
):
    """
    Reconstruct Groww turnover from the actual
    trade-level table.

    Observed Groww methodology in supplied reports:

        turnover =
        SUM(ABS(trade-level Realized P&L))

    Also independently sums signed trade-level
    P&L so it can be reconciled against the
    Realised P&L summary.

    The P&L column is detected by header rather
    than assuming column I.
    """

    trade_heading = (
        find_exact(
            sheet,
            "Realised trades (trade level)",
        )
        or
        find_exact(
            sheet,
            "Realized trades (trade level)",
        )
    )

    if not trade_heading:
        raise GrowwParseError(
            "Groww trade-level section "
            "was not found."
        )

    turnover = Decimal("0")
    signed_pnl = Decimal("0")

    trade_count = 0

    pnl_column = None

    for row in range(
        trade_heading.row + 1,
        sheet.max_row + 1,
    ):

        first_value = normalize_text(
            sheet.cell(
                row=row,
                column=1,
            ).value
        )

        first_lower = (
            first_value.casefold()
        )

        # -----------------------------------------
        # End of report
        # -----------------------------------------

        if (
            first_lower
            == "disclaimer:"
            or first_lower
            == "disclaimer"
        ):
            break

        # -----------------------------------------
        # Detect table header dynamically
        # -----------------------------------------

        detected_pnl_column = None

        for column in range(
            1,
            sheet.max_column + 1,
        ):

            header = normalize_text(
                sheet.cell(
                    row=row,
                    column=column,
                ).value
            ).casefold()

            if header in {
                "realized p&l",
                "realised p&l",
            }:

                detected_pnl_column = (
                    column
                )

                break

        if (
            detected_pnl_column
            is not None
        ):

            pnl_column = (
                detected_pnl_column
            )

            continue

        # -----------------------------------------
        # Section labels
        # -----------------------------------------

        if first_lower in {
            "",
            "futures",
            "options",
            "scrip name",
        }:
            continue

        if pnl_column is None:
            continue

        # -----------------------------------------
        # Never count summary/total rows
        # -----------------------------------------

        if first_lower == "total":
            continue

        raw_pnl = sheet.cell(
            row=row,
            column=pnl_column,
        ).value

        pnl = try_decimal(
            raw_pnl
        )

        if pnl is None:
            continue

        signed_pnl += pnl

        turnover += abs(
            pnl
        )

        trade_count += 1

    if trade_count == 0:
        raise GrowwParseError(
            "No Groww trade-level rows "
            "were found."
        )

    return (
        turnover,
        signed_pnl,
    )


# =========================================================
# MAIN F&O / COMMODITY PARSER
# =========================================================


def parse_groww_report(
    filepath: str | Path,
    report_type: str,
) -> GrowwParseResult:

    report_type = (
        report_type.upper()
    )

    if report_type not in {
        "FNO",
        "COMMODITY",
    }:

        raise GrowwParseError(
            "This parser supports only "
            "Groww FNO and COMMODITY reports."
        )

    # -----------------------------------------------------
    # Workbook loading
    # -----------------------------------------------------

    try:

        workbook = load_workbook(
            filename=filepath,
            data_only=True,
            read_only=False,
        )

    except (
        InvalidFileException,
        BadZipFile,
        OSError,
        ValueError,
    ) as exc:

        raise GrowwParseError(
            "The uploaded file is not a "
            "readable Excel workbook."
        ) from exc

    # -----------------------------------------------------
    # Identify correct worksheet from selected type
    # -----------------------------------------------------

    (
        sheet,
        title,
    ) = find_title_sheet(
        workbook,
        report_type,
    )

    # -----------------------------------------------------
    # Identity
    # -----------------------------------------------------

    (
        client_name,
        client_code,
    ) = extract_identity(
        sheet
    )

    # -----------------------------------------------------
    # Period
    # -----------------------------------------------------

    (
        period_start,
        period_end,
    ) = extract_period(
        title
    )

    # -----------------------------------------------------
    # Summary
    # -----------------------------------------------------

    realised_pnl = (
        find_section_total(
            sheet,
            "Realised P&L",
        )
    )

    charges = (
        find_section_total(
            sheet,
            "Charges",
        )
    )

    reported_turnover = (
        find_optional_turnover(
            sheet
        )
    )

    # -----------------------------------------------------
    # Independent trade-level calculation
    # -----------------------------------------------------

    (
        calculated_turnover,
        trade_level_pnl,
    ) = (
        calculate_trade_level_turnover(
            sheet
        )
    )

    # -----------------------------------------------------
    # Groww business P&L rule
    # -----------------------------------------------------

    final_net_pnl = (
        realised_pnl.value
        -
        charges.value
    )

    # -----------------------------------------------------
    # Turnover choice
    # -----------------------------------------------------

    if reported_turnover:

        turnover_source = (
            "BROKER_REPORTED"
        )

    else:

        turnover_source = (
            "CALCULATED_FROM_TRADE_LEVEL"
        )

    return GrowwParseResult(
        report_type=report_type,

        client_name=client_name,
        client_code=client_code,

        period_start=period_start,
        period_end=period_end,

        realised_pnl=realised_pnl,

        charges=charges,

        reported_turnover=(
            reported_turnover
        ),

        calculated_turnover=(
            calculated_turnover
        ),

        turnover_source=(
            turnover_source
        ),

        trade_level_pnl=(
            trade_level_pnl
        ),

        final_net_pnl=(
            final_net_pnl
        ),
    )