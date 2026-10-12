from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
import re

from openpyxl import load_workbook

from app.brokers.groww.parser import (
    GrowwParseError,
    SourceValue,
    normalize_text,
    to_decimal,
)


@dataclass
class GrowwMutualFundParseResult:
    report_type: str

    client_name: str
    pan: str

    period_start: date
    period_end: date

    stcg: SourceValue
    ltcg: SourceValue

    sale_consideration: Decimal

    transaction_stcg: Decimal
    transaction_ltcg: Decimal

    redemption_count: int


# =========================================================
# FIND REAL MF SHEET
# =========================================================

def find_mutual_fund_sheet(
    workbook,
):
    """
    Actual supplied Groww report uses:

        Capital gains

    We do not hard-code only that name.
    Find the sheet from its report title / headers.
    """

    for sheet in workbook.worksheets:

        for row in range(
            1,
            min(sheet.max_row, 25) + 1,
        ):

            for column in range(
                1,
                min(sheet.max_column, 6) + 1,
            ):

                value = normalize_text(
                    sheet.cell(
                        row=row,
                        column=column,
                    ).value
                )

                if (
                    "CAPITAL GAINS STATEMENT FOR FY"
                    in value.upper()
                ):
                    return sheet

    raise GrowwParseError(
        "Could not locate a supported "
        "Groww Mutual Funds Capital Gains "
        "sheet."
    )


# =========================================================
# PERSONAL DETAILS
# =========================================================

def find_label_value(
    sheet,
    label: str,
):

    target = label.casefold()

    for row in range(
        1,
        min(sheet.max_row, 20) + 1,
    ):

        for column in range(
            1,
            sheet.max_column + 1,
        ):

            value = normalize_text(
                sheet.cell(
                    row=row,
                    column=column,
                ).value
            )

            if value.casefold() == target:

                result = sheet.cell(
                    row=row,
                    column=column + 1,
                ).value

                return normalize_text(
                    result
                )

    return None


# =========================================================
# FY
# =========================================================

def extract_financial_year(
    sheet,
):

    title = None

    for row in range(
        1,
        min(sheet.max_row, 25) + 1,
    ):

        for column in range(
            1,
            min(sheet.max_column, 6) + 1,
        ):

            value = normalize_text(
                sheet.cell(
                    row=row,
                    column=column,
                ).value
            )

            if (
                "CAPITAL GAINS STATEMENT FOR FY"
                in value.upper()
            ):
                title = value
                break

        if title:
            break

    if not title:
        raise GrowwParseError(
            "Could not determine Mutual "
            "Funds financial year."
        )

    match = re.search(
        r"FY\s+(\d{4})-(\d{2})",
        title,
        re.IGNORECASE,
    )

    if not match:
        raise GrowwParseError(
            "Invalid Groww Mutual Funds "
            "financial year format."
        )

    start_year = int(
        match.group(1)
    )

    expected_end = (
        start_year + 1
    ) % 100

    actual_end = int(
        match.group(2)
    )

    if (
        actual_end
        != expected_end
    ):
        raise GrowwParseError(
            "Groww Mutual Funds financial "
            "year is inconsistent."
        )

    return (
        date(
            start_year,
            4,
            1,
        ),
        date(
            start_year + 1,
            3,
            31,
        ),
    )


# =========================================================
# SUMMARY
# =========================================================

def extract_capital_gain_summary(
    sheet,
):
    """
    Actual summary structure:

    Asset Class / Category
    Taxable Short Term
    Taxable Long Term

    Equity
    Debt ...
    Debt ...

    We sum all taxable categories.
    """

    header_row = None
    category_column = None
    stcg_column = None
    ltcg_column = None

    for row in range(
        1,
        min(sheet.max_row, 25) + 1,
    ):

        for column in range(
            1,
            sheet.max_column + 1,
        ):

            value = normalize_text(
                sheet.cell(
                    row=row,
                    column=column,
                ).value
            )

            if (
                value.casefold()
                == "asset class / category"
            ):

                header_row = row
                category_column = column

                break

        if header_row:
            break

    if not header_row:
        raise GrowwParseError(
            "Mutual Fund capital gain "
            "summary was not found."
        )

    for column in range(
        category_column,
        sheet.max_column + 1,
    ):

        header = normalize_text(
            sheet.cell(
                row=header_row,
                column=column,
            ).value
        ).casefold()

        if header == "taxable short term":
            stcg_column = column

        elif header == "taxable long term":
            ltcg_column = column

    if stcg_column is None:
        raise GrowwParseError(
            "Taxable Short Term column "
            "was not found."
        )

    if ltcg_column is None:
        raise GrowwParseError(
            "Taxable Long Term column "
            "was not found."
        )

    total_stcg = Decimal("0")
    total_ltcg = Decimal("0")

    row = header_row + 1

    first_stcg_cell = None
    first_ltcg_cell = None

    found_rows = 0

    while row <= sheet.max_row:

        category = normalize_text(
            sheet.cell(
                row=row,
                column=category_column,
            ).value
        )

        if not category:
            break

        raw_stcg = sheet.cell(
            row=row,
            column=stcg_column,
        ).value

        raw_ltcg = sheet.cell(
            row=row,
            column=ltcg_column,
        ).value

        if (
            raw_stcg is None
            and raw_ltcg is None
        ):
            break

        if first_stcg_cell is None:
            first_stcg_cell = sheet.cell(
                row=row,
                column=stcg_column,
            )

            first_ltcg_cell = sheet.cell(
                row=row,
                column=ltcg_column,
            )

        total_stcg += to_decimal(
            raw_stcg or 0
        )

        total_ltcg += to_decimal(
            raw_ltcg or 0
        )

        found_rows += 1
        row += 1

    if found_rows == 0:
        raise GrowwParseError(
            "No Mutual Fund capital gain "
            "summary rows were found."
        )

    stcg = SourceValue(
        value=total_stcg,
        sheet=sheet.title,
        cell=(
            first_stcg_cell.coordinate
            if first_stcg_cell
            else ""
        ),
        label=(
            "Total Taxable Short Term"
        ),
    )

    ltcg = SourceValue(
        value=total_ltcg,
        sheet=sheet.title,
        cell=(
            first_ltcg_cell.coordinate
            if first_ltcg_cell
            else ""
        ),
        label=(
            "Total Taxable Long Term"
        ),
    )

    return (
        stcg,
        ltcg,
    )


# =========================================================
# TRANSACTION TABLE
# =========================================================

def calculate_redemptions(
    sheet,
):

    header_row = None

    matched_quantity_column = None
    redeem_price_column = None
    stcg_column = None
    ltcg_column = None

    for row in range(
        1,
        sheet.max_row + 1,
    ):

        headers = {}

        for column in range(
            1,
            sheet.max_column + 1,
        ):

            value = normalize_text(
                sheet.cell(
                    row=row,
                    column=column,
                ).value
            )

            if value:
                headers[
                    value.casefold()
                ] = column

        if (
            "matched quantity"
            in headers
            and
            "redeem price"
            in headers
        ):

            header_row = row

            matched_quantity_column = (
                headers[
                    "matched quantity"
                ]
            )

            redeem_price_column = (
                headers[
                    "redeem price"
                ]
            )

            stcg_column = headers.get(
                "short term-capital gain"
            )

            ltcg_column = headers.get(
                "long term-capital gain"
            )

            break

    if header_row is None:
        raise GrowwParseError(
            "Mutual Fund redemption "
            "transaction table was not found."
        )

    if stcg_column is None:
        raise GrowwParseError(
            "Short Term-Capital Gain "
            "column was not found."
        )

    if ltcg_column is None:
        raise GrowwParseError(
            "Long Term-Capital Gain "
            "column was not found."
        )

    sale_consideration = (
        Decimal("0")
    )

    transaction_stcg = Decimal("0")
    transaction_ltcg = Decimal("0")

    count = 0

    row = header_row + 1

    while row <= sheet.max_row:

        scheme_name = normalize_text(
            sheet.cell(
                row=row,
                column=1,
            ).value
        )

        matched_quantity = sheet.cell(
            row=row,
            column=(
                matched_quantity_column
            ),
        ).value

        redeem_price = sheet.cell(
            row=row,
            column=redeem_price_column,
        ).value

        # End of transaction table
        if (
            not scheme_name
            or matched_quantity is None
            or redeem_price is None
        ):
            break

        quantity = to_decimal(
            matched_quantity
        )

        price = to_decimal(
            redeem_price
        )

        # IMPORTANT:
        # Transaction-by-transaction.
        sale_consideration += (
            price
            *
            quantity
        )

        transaction_stcg += to_decimal(
            sheet.cell(
                row=row,
                column=stcg_column,
            ).value
            or 0
        )

        transaction_ltcg += to_decimal(
            sheet.cell(
                row=row,
                column=ltcg_column,
            ).value
            or 0
        )

        count += 1
        row += 1

    if count == 0:
        raise GrowwParseError(
            "No Mutual Fund redemption "
            "transactions were found."
        )

    return (
        sale_consideration,
        transaction_stcg,
        transaction_ltcg,
        count,
    )


# =========================================================
# PARSER
# =========================================================

def parse_groww_mutual_funds(
    filepath: str | Path,
) -> GrowwMutualFundParseResult:

    workbook = load_workbook(
        filename=filepath,
        data_only=True,
        read_only=False,
    )

    sheet = find_mutual_fund_sheet(
        workbook
    )

    client_name = find_label_value(
        sheet,
        "Name",
    )

    pan = find_label_value(
        sheet,
        "PAN",
    )

    if not client_name:
        raise GrowwParseError(
            "Mutual Fund client name "
            "was not found."
        )

    if not pan:
        raise GrowwParseError(
            "Mutual Fund PAN "
            "was not found."
        )

    (
        period_start,
        period_end,
    ) = extract_financial_year(
        sheet
    )

    (
        stcg,
        ltcg,
    ) = extract_capital_gain_summary(
        sheet
    )

    (
        sale_consideration,
        transaction_stcg,
        transaction_ltcg,
        redemption_count,
    ) = calculate_redemptions(
        sheet
    )

    return GrowwMutualFundParseResult(
        report_type="MUTUAL_FUNDS",

        client_name=client_name,
        pan=pan,

        period_start=period_start,
        period_end=period_end,

        stcg=stcg,
        ltcg=ltcg,

        sale_consideration=(
            sale_consideration
        ),

        transaction_stcg=(
            transaction_stcg
        ),

        transaction_ltcg=(
            transaction_ltcg
        ),

        redemption_count=(
            redemption_count
        ),
    )